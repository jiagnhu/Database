import io,os,re,socket,struct,subprocess,tempfile,zipfile
from pathlib import Path
from transport import ProcessingError

def scan(data):
    host=os.getenv('CLAMAV_HOST','clamav');port=int(os.getenv('CLAMAV_PORT','3310'))
    try:
        with socket.create_connection((host,port),timeout=20) as s:
            s.settimeout(90);s.sendall(b'zINSTREAM\0')
            for i in range(0,len(data),65536):
                chunk=data[i:i+65536];s.sendall(struct.pack('!I',len(chunk))+chunk)
            s.sendall(struct.pack('!I',0));reply=b''
            while b'\0' not in reply and len(reply)<4096:
                part=s.recv(1024)
                if not part:break
                reply+=part
            if b'FOUND' in reply:raise ProcessingError('VIRUS_DETECTED')
            if b'OK' not in reply:raise ProcessingError('VIRUS_SCAN_FAILED')
    except OSError:raise ProcessingError('VIRUS_SCAN_UNAVAILABLE') from None

def check_office(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            entries=z.infolist()
            if len(entries)>3000 or sum(x.file_size for x in entries)>150*1024*1024:raise ProcessingError('ARCHIVE_TOO_LARGE')
            for entry in entries:
                if entry.file_size>30*1024*1024 or entry.file_size/max(1,entry.compress_size)>300:raise ProcessingError('ARCHIVE_TOO_LARGE')
                if '..' in Path(entry.filename).parts or entry.filename.startswith(('/','\\')):raise ProcessingError('ARCHIVE_PATH')
                if 'vbaproject' in entry.filename.lower():raise ProcessingError('MACRO_REJECTED')
                if entry.filename.lower().endswith(('.xml','.rels')):
                    content=z.read(entry)
                    if b'<!DOCTYPE' in content or b'<!ENTITY' in content:raise ProcessingError('XML_ENTITY_REJECTED')
    except zipfile.BadZipFile:raise ProcessingError('INVALID_OFFICE_FILE') from None

def sentences(text):
    text=re.sub(r'\s+',' ',text).strip()
    if not text:return []
    parts=re.split(r'(?<=[.!?。！？])\s+(?=[A-Z\u4e00-\u9fff])',text)
    return [piece[i:i+1100] for piece in parts for i in range(0,len(piece),1100) if piece[i:i+1100]]

def extract(data,kind,heartbeat=lambda:None):
    pages=[];max_pages=int(os.getenv('WORKER_MAX_PAGES','200'));max_chars=int(os.getenv('WORKER_MAX_CHARACTERS','150000'));chars=0
    def append(label,texts,ocr=False):
        nonlocal chars
        if len(pages)>=max_pages:raise ProcessingError('PAGE_LIMIT')
        texts=[re.sub(r'\s+',' ',t).strip() for t in texts if t.strip()]
        chars+=sum(map(len,texts))
        if chars>max_chars:raise ProcessingError('TEXT_LIMIT')
        pairs=[]
        for i,t in enumerate(texts):
            for j,s in enumerate(sentences(t)):
                chinese=len(re.findall(r'[\u4e00-\u9fff]',s))>len(s)*.35
                pairs.append({'id':f'p{len(pages)+1}b{i}s{j}','en':'' if chinese else s,'zh':s if chinese else '', 'kind':'text'})
        pages.append({'id':f's{len(pages)+1}','title':texts[0][:80] if texts else label,'source':label+(' · OCR' if ocr else ''),'pairs':pairs})
        heartbeat()
    if kind=='pdf':
        from pypdf import PdfReader
        pdf=PdfReader(io.BytesIO(data),strict=False)
        if pdf.is_encrypted and not pdf.decrypt(''):raise ProcessingError('ENCRYPTED_PDF')
        if len(pdf.pages)>max_pages:raise ProcessingError('PAGE_LIMIT')
        rendered=None
        try:
            for i,page in enumerate(pdf.pages):
                text=page.extract_text() or '';ocr=len(re.sub(r'\s','',text))<30
                if ocr:
                    import pypdfium2 as pdfium
                    rendered=rendered or pdfium.PdfDocument(data)
                    p=rendered[i]
                    if p.get_width()*p.get_height()*4>30000000:raise ProcessingError('OCR_PAGE_TOO_LARGE')
                    bitmap=p.render(scale=2);image=bitmap.to_pil()
                    with tempfile.TemporaryDirectory(prefix='cpv-ocr-') as tmp:
                        name=str(Path(tmp)/'page.png');image.save(name)
                        proc=subprocess.run(['tesseract',name,'stdout','-l','eng+chi_sim'],capture_output=True,timeout=60,check=False)
                        if proc.returncode:raise ProcessingError('OCR_FAILED')
                        text=proc.stdout.decode('utf8',errors='replace')
                    image.close();bitmap.close();p.close()
                append(f'PDF 第 {i+1} 页',text.splitlines(),ocr)
        finally:
            if rendered:rendered.close()
    elif kind=='docx':
        check_office(data)
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        doc=Document(io.BytesIO(data));texts=[]
        for child in doc.element.body:
            if child.tag.endswith('}p'):
                t=Paragraph(child,doc).text
                if t.strip():texts.append(t)
            elif child.tag.endswith('}tbl'):
                for row in Table(child,doc).rows:texts.append(' | '.join(c.text for c in row.cells))
        for i in range(0,len(texts),20):append(f'原文段落 {i+1}–{min(i+20,len(texts))}',texts[i:i+20])
    elif kind=='pptx':
        check_office(data)
        from pptx import Presentation
        prs=Presentation(io.BytesIO(data))
        if len(prs.slides)>max_pages:raise ProcessingError('PAGE_LIMIT')
        def texts(shapes):
            for s in shapes:
                if hasattr(s,'shapes'):yield from texts(s.shapes)
                if getattr(s,'has_text_frame',False):
                    for p in s.text_frame.paragraphs:yield p.text
                if getattr(s,'has_table',False):
                    for row in s.table.rows:yield ' | '.join(c.text for c in row.cells)
        for i,slide in enumerate(prs.slides):append(f'原课件第 {i+1} 页',list(texts(slide.shapes)))
    else:raise ProcessingError('UNSUPPORTED_FORMAT')
    if not any(p['pairs'] for p in pages):raise ProcessingError('NO_TEXT_EXTRACTED')
    return pages
